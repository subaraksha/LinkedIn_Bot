"""Resource-bounded resume text parser. Run only as an isolated child process."""

import json
import resource
import sys
import zipfile
from pathlib import Path

MAX_INPUT = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024
MAX_ENTRIES = 2000
MAX_PAGES = 300
MAX_TEXT = 500_000
MEMORY_LIMIT = 512 * 1024 * 1024


class ParseError(ValueError):
    pass


def check_docx_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ENTRIES:
                raise ParseError("DOCX has too many archive entries")
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(archive.namelist()):
                raise ParseError("File is not a DOCX document")
            total = 0
            for entry in entries:
                name = entry.filename
                if (name.startswith("/") or "\\" in name or
                        any(part in {"", ".", ".."} for part in name.rstrip("/").split("/"))):
                    raise ParseError("DOCX has unsafe archive paths")
                if entry.flag_bits & 1:
                    raise ParseError("Encrypted DOCX is not supported")
                total += entry.file_size
                if total > MAX_EXPANDED:
                    raise ParseError("DOCX expanded content is too large")
    except (zipfile.BadZipFile, OSError) as exc:
        raise ParseError("DOCX could not be read") from exc


def extract(path: Path, kind: str) -> list[dict[str, str]]:
    if path.stat().st_size > MAX_INPUT:
        raise ParseError("Resume exceeds the 10 MiB limit")
    if kind == "txt":
        raw = path.read_bytes()
        if b"\x00" in raw:
            raise ParseError("File is not plain text")
        try:
            value = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ParseError("Plain text must be UTF-8") from exc
        sections = [{"label": "Document", "text": value}]
    elif kind == "pdf":
        if not path.open("rb").read(5).startswith(b"%PDF-"):
            raise ParseError("File is not a PDF")
        from pypdf import PdfReader
        try:
            reader = PdfReader(str(path), strict=True)
            if reader.is_encrypted:
                raise ParseError("Encrypted PDF is not supported")
            if len(reader.pages) > MAX_PAGES:
                raise ParseError("PDF has too many pages")
            sections = [{"label": f"Page {index + 1}", "text": page.extract_text() or ""}
                        for index, page in enumerate(reader.pages)]
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError("PDF could not be read") from exc
    elif kind == "docx":
        check_docx_archive(path)
        from docx import Document
        try:
            document = Document(str(path))
            sections = [{"label": f"Paragraph {index + 1}", "text": paragraph.text}
                        for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip()]
            for table_index, table in enumerate(document.tables):
                for row_index, row in enumerate(table.rows):
                    text = " | ".join(cell.text for cell in row.cells)
                    if text.strip():
                        sections.append({"label": f"Table {table_index + 1}, row {row_index + 1}", "text": text})
        except Exception as exc:
            raise ParseError("DOCX could not be read") from exc
    else:
        raise ParseError("Unsupported resume format")
    cleaned = []
    total = 0
    for section in sections:
        value = section["text"].strip().replace("\r\n", "\n")
        if value:
            total += len(value)
            if total > MAX_TEXT:
                raise ParseError("Extracted resume text is too long")
            cleaned.append({"label": section["label"], "text": value})
    if total < 20:
        raise ParseError("No useful text found; scanned files need pasted text")
    return cleaned


def main() -> None:
    try:
        if sys.platform != "darwin":
            resource.setrlimit(resource.RLIMIT_AS, (MEMORY_LIMIT, MEMORY_LIMIT))
        resource.setrlimit(resource.RLIMIT_CPU, (30, 30))
        sections = extract(Path(sys.argv[1]), sys.argv[2])
        print(json.dumps({"sections": sections}, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc) if isinstance(exc, ParseError) else "Resume parsing failed"}))
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
